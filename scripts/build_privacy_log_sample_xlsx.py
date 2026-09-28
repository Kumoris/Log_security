"""Build privacy_log_sample_500.xlsx from workbook_payload.json (openpyxl port
of build_privacy_log_sample_workbook.mjs, used when @oai/artifact-tool is absent)."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
payload = json.loads((ROOT / "data/res/privacy_log_sample_500/workbook_payload.json").read_text("utf-8"))
out_path = ROOT / "outputs/privacy_log_sample_500/privacy_log_sample_500.xlsx"
out_path.parent.mkdir(parents=True, exist_ok=True)

HEADER_FILL = PatternFill("solid", fgColor="F1F3F4")
HEADER_FONT = Font(bold=True, color="202124")
RED_FILL = PatternFill("solid", fgColor="FCE8E6")
RED_FONT = Font(bold=True, color="A50E0E")
AMBER_FILL = PatternFill("solid", fgColor="FEF7E0")
AMBER_FONT = Font(bold=True, color="8A5A00")
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
wb.remove(wb.active)

TAX_DROPDOWN = ["IL-At", "IL-Pa", "IL-Lv", "SS-Cr", "SS-Cf", "SS-Ur", "RM-Ms", "RM-Ft",
                "EE-Ex", "EE-St", "LINK-Id", "OUT"]
PR_DROPDOWN = ["PR-CRED", "PR-USER", "PR-LINK-ID", "PR-CONFIG", "PR-RAW-DUMP", "PR-ERROR",
               "PR-REDACTION", "PR-LOG-INTEGRITY", "PR-NONE"]
NEWCAT_DROPDOWN = ["none", "NEW-RAW-STRUCTURE-DUMP", "NEW-LINKABLE-IDENTIFIER", "NEW-INFRA-TOPOLOGY",
                   "NEW-LLM-IO", "NEW-OBSERVABILITY-OVERCOLLECTION", "NEW-TEST-DEMO-SECRET", "other"]

type_rank = {"SS-Cr": 1, "SS-Ur": 2, "SS-Cf": 3, "RM-Ms": 4, "RM-Ft": 5, "EE-Ex": 6,
             "EE-St": 7, "LINK-Id": 8, "IL-At": 9, "IL-Pa": 10, "IL-Lv": 11,
             "OUT-NonLog": 98, "OUT-Low": 99}


def write_sheet(name, rows, freeze_cols=0, widths=None):
    ws = wb.create_sheet(name)
    for r in rows:
        ws.append(["" if c is None else c for c in r])
    if rows:
        for ci in range(1, len(rows[0]) + 1):
            cell = ws.cell(row=1, column=ci)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = ws.cell(row=2, column=freeze_cols + 1)
    if widths:
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
    return ws


def records_to_rows(records):
    if not records:
        return [[]]
    headers = list(records[0].keys())
    return [headers] + [[rec.get(h, "") for h in headers] for rec in records]


def add_dv(ws, col_letter, options, n):
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(options), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{n}")


def row_problem_type(row):
    if row.get("privacy_security_issue_candidate") != "True":
        return "NO-ISSUE - 未发现隐私/安全候选问题"
    pp = row.get("primary_pattern", "")
    if pp and pp not in ("OUT-Low", "OUT-NonLog"):
        return f"{pp} - {row.get('taxonomy_label','')}"
    return f"{row.get('primary_privacy_leak_type','')} - {row.get('primary_privacy_leak_label','')}"


# ---- Log_List_By_Type ----
sample = list(payload["sample_rows"])
sample.sort(key=lambda r: (
    0 if r.get("privacy_security_issue_candidate") == "True" else 1,
    type_rank.get(r.get("primary_pattern"), 50),
    int(r.get("sample_order") or 0),
))
list_records = []
for idx, row in enumerate(sample, 1):
    pp = row.get("primary_pattern", "")
    list_records.append({
        "list_no": idx,
        "original_sample_no": row.get("sample_order"),
        "has_privacy_security_issue": "是" if row.get("privacy_security_issue_candidate") == "True" else "否",
        "problem_type": row_problem_type(row),
        "paper_taxonomy_pattern": "OUT" if pp in ("OUT-Low", "OUT-NonLog") else pp,
        "paper_taxonomy_label": row.get("taxonomy_label"),
        "privacy_leak_type": row.get("primary_privacy_leak_type"),
        "privacy_leak_label": row.get("primary_privacy_leak_label"),
        "new_category_candidate": row.get("new_error_labels") if row.get("new_error_candidate") == "True" else "",
        "log": row.get("log_text_for_review"),
        "reason": row.get("privacy_classification_reason") or row.get("classification_reason"),
        "dataset": row.get("dataset"), "repo": row.get("repo"), "file": row.get("file"),
        "line": row.get("line"), "attribution": row.get("attribution"), "risk_level": row.get("risk_level"),
        "source_locator": row.get("source_locator"),
        "manual_final_yes_no": "", "manual_final_type": "", "manual_new_category": "", "manual_notes": "",
    })
list_rows = records_to_rows(list_records)
n = len(list_rows)
ll = write_sheet("Log_List_By_Type", list_rows, freeze_cols=4, widths={
    "A": 8, "B": 12, "C": 14, "D": 44, "E": 16, "F": 22, "G": 16, "H": 22,
    "I": 34, "J": 70, "K": 48, "L": 16, "M": 26, "N": 40, "O": 8, "P": 16, "Q": 12, "R": 40,
    "S": 18, "T": 20, "U": 26, "V": 24})
for col in ["D", "I", "J", "K", "N", "R"]:
    for r in range(2, n + 1):
        ll[f"{col}{r}"].alignment = WRAP
ll.conditional_formatting.add(f"C2:C{n}", FormulaRule(formula=['EXACT(C2,"是")'], fill=RED_FILL, font=RED_FONT))
ll.conditional_formatting.add(f"I2:I{n}", FormulaRule(formula=['ISNUMBER(SEARCH("NEW-",I2))'], fill=AMBER_FILL, font=AMBER_FONT))
add_dv(ll, "S", ["是", "否", "不确定", "需要上下文"], n)
add_dv(ll, "T", TAX_DROPDOWN + PR_DROPDOWN[:-1] + ["NO-ISSUE", "OTHER"], n)
add_dv(ll, "U", NEWCAT_DROPDOWN, n)

# ---- Type_Index ----
type_index = [["problem_type", "rows"]]
counts = {}
for rec in list_records:
    counts[rec["problem_type"]] = counts.get(rec["problem_type"], 0) + 1
for pt, c in counts.items():
    type_index.append([pt, c])
write_sheet("Type_Index", type_index, widths={"A": 64, "B": 12})

# ---- Summary ----
s = payload["summary"]
summary_rows = [
    ["Privacy Log Sample 500", "", ""],
    ["Metric", "Value", "Notes"],
    ["Source rows", s["method"]["source_rows"], "All extracted log candidates across AIDev, SWE-chat, DevGPT"],
    ["Sample rows", s["overall"]["sample_rows"], f"Uniform random sample; seed {s['method']['random_seed']}"],
    ["Privacy/security candidate rows", s["overall"]["privacy_security_candidates"], "Candidate labels only; manual review required"],
    ["Security taxonomy candidate rows", s["overall"]["security_taxonomy_candidates"], "Paper taxonomy heuristic (improved rules)"],
    ["New error candidate rows", s["overall"]["new_error_candidate_rows"], "Potential innovation labels"],
    ["Privacy/security candidate rate", f"{s['overall']['privacy_security_candidate_rate_percent']}%", ""],
    ["New error candidate rate", f"{s['overall']['new_error_candidate_rate_percent']}%", ""],
    ["Paper basis", s["method"]["paper"], s["method"]["paper_taxonomy"]],
    ["Caveat", s["method"]["caveat"], ""],
    ["", "", ""],
    ["By Dataset", "", ""],
    ["Dataset", "Sample Rows", "Privacy/Security Candidates", "Security Taxonomy Candidates", "New Error Candidate Rows"],
    *[[r["dataset"], r["sample_rows"], r["privacy_security_candidates"], r["security_taxonomy_candidates"], r["new_error_candidate_rows"]] for r in s["by_dataset"]],
    ["", "", "", "", ""],
    ["By Privacy Leak Type", "", ""],
    ["Code", "Name", "Rows"],
    *[[r["label"], r["name"], r["rows"]] for r in s["by_privacy_type"]],
    ["", "", ""],
    ["By New Error Candidate", "", ""],
    ["Code", "Name", "Rows"],
    *[[r["label"], r["name"], r["rows"]] for r in s["by_new_error_candidate"]],
    ["", "", ""],
    ["By Paper Taxonomy Pattern", "", ""],
    ["Pattern", "Name", "Rows"],
    *[[r["pattern"], r["name"], r["rows"]] for r in s["by_taxonomy_pattern"]],
]
ss = write_sheet("Summary", summary_rows, widths={"A": 34, "B": 24, "C": 72, "D": 26, "E": 26})
ss.merge_cells("A1:C1")
ss["A1"].font = Font(bold=True, size=16, color="174EA6")
ss["A1"].fill = PatternFill("solid", fgColor="E8F0FE")

# ---- Review_500 ----
review_rows = records_to_rows(payload["sample_rows"])
headers = review_rows[0]
nr = len(review_rows)
rv = write_sheet("Review_500", review_rows, freeze_cols=8)
def col_of(name):
    return get_column_letter(headers.index(name) + 1)
add_dv(rv, col_of("manual_review_status"), ["not reviewed", "confirmed issue", "false positive", "uncertain", "needs context"], nr)
add_dv(rv, col_of("human_final_issue"), ["yes", "no", "uncertain"], nr)
add_dv(rv, col_of("human_final_taxonomy_pattern"), TAX_DROPDOWN, nr)
add_dv(rv, col_of("human_final_privacy_type"), PR_DROPDOWN, nr)
add_dv(rv, col_of("human_new_category"), NEWCAT_DROPDOWN, nr)
pc = col_of("privacy_security_issue_candidate")
rv.conditional_formatting.add(f"{pc}2:{pc}{nr}", FormulaRule(formula=[f'EXACT({pc}2,"True")'], fill=RED_FILL, font=RED_FONT))
nc = col_of("new_error_candidate")
rv.conditional_formatting.add(f"{nc}2:{nc}{nr}", FormulaRule(formula=[f'EXACT({nc}2,"True")'], fill=AMBER_FILL, font=AMBER_FONT))
for c in ["O", "AJ"]:
    rv.column_dimensions[c].width = 60

# ---- Paper_Taxonomy / New_Category_Candidates / Privacy_Type_Legend ----
write_sheet("Paper_Taxonomy", records_to_rows(payload["taxonomy_rows"]), widths={"A": 14, "B": 30, "C": 12, "D": 30, "E": 76})
write_sheet("New_Category_Candidates", records_to_rows(payload["new_category_rows"]), widths={"A": 34, "B": 40, "C": 12, "D": 84})
write_sheet("Privacy_Type_Legend", records_to_rows(payload["privacy_type_rows"]), widths={"A": 22, "B": 50})

wb.save(out_path)
print(f"wrote {out_path}")
print("sheets:", wb.sheetnames)
